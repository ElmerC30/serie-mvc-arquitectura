"""Genera Respuestas_Serie_MVC.docx con las respuestas teoricas, el analisis y el registro de
resultados esperados vs. obtenidos de los ejercicios 0 a 8.

Requiere haber corrido antes: python scripts/verificar_todo.py (produce evidencias/).
Uso: python scripts/generar_docx.py
"""
import json
import pathlib
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

RAIZ = pathlib.Path(__file__).resolve().parent.parent
EVID = RAIZ / "evidencias"
SALIDA = RAIZ / "Respuestas_Serie_MVC.docx"
ANCHO = Cm(16.5)

COLOR = {"M": "C6EFCE", "V": "DDEBF7", "C": "FCE4D6", "MIX": "FFEB9C", "-": "FFFFFF"}
NOMBRE = {"M": "Modelo", "V": "Vista", "C": "Controlador", "-": "(estructura)"}


# ============================================================================ utilidades de formato

def sombrear(celda, color):
    celda._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color}"/>'))


def ancho_celdas(tabla, anchos_cm):
    tabla.autofit = False
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos_cm):
            celda.width = Cm(ancho)


def parrafo(doc, texto="", negrita=False, cursiva=False, tam=None, alinear=None, espacio_despues=None):
    p = doc.add_paragraph()
    if texto:
        agregar_texto(p, texto, negrita, cursiva, tam)
    if alinear:
        p.alignment = alinear
    if espacio_despues is not None:
        p.paragraph_format.space_after = Pt(espacio_despues)
    return p


# Acentos de la prosa. El codigo fuente y los identificadores no se tocan: solo se aplica a texto normal.
_ACENTOS = {w.split(":")[0]: w.split(":")[1] for w in """
decision:decisión decisiones:decisiones informacion:información presentacion:presentación validacion:validación
notificacion:notificación aplicacion:aplicación implementacion:implementación implementaciones:implementaciones
operacion:operación excepcion:excepción conversion:conversión coordinacion:coordinación linea:línea lineas:líneas
numero:número numeros:números metodo:método metodos:métodos codigo:código analisis:análisis patron:patrón
ultimo:último ultima:última unica:única unico:único logica:lógica grafica:gráfica graficas:gráficas
practica:práctica tambien:también despues:después ademas:además asi:así aqui:aquí segun:según ningun:ningún
ninguna:ninguna algun:algún dia:día dias:días mas:más fragil:frágil valido:válido validos:válidos
invalido:inválido invalidos:inválidos invalida:inválida invalidas:inválidas tecnologia:tecnología
tecnologias:tecnologías deberia:debería seria:sería tendria:tendría tendrian:tendrían habria:habría
necesitarian:necesitarían dependeria:dependería tocaria:tocaría revision:revisión correccion:corrección
funcion:función sesion:sesión peticion:petición demostracion:demostración simulacion:simulación
direccion:dirección separacion:separación transcripcion:transcripción inversion:inversión
comparacion:comparación justificacion:justificación verificacion:verificación verificaciones:verificaciones
exportacion:exportación navegacion:navegación confirmacion:confirmación representacion:representación
actualizacion:actualización suscripcion:suscripción Ingenieria:Ingeniería TECNOLOGICO:TECNOLÓGICO
proposito:propósito clasico:clásico clasica:clásica titulo:título minimo:mínimo minimos:mínimos maximo:máximo
pequeno:pequeño diseno:diseño util:útil estatico:estático estatica:estática estaticas:estáticas
estaticos:estáticos tipicas:típicas tipico:típico vacio:vacío vacia:vacía leccion:lección estandar:estándar
comun:común nucleo:núcleo categoria:categoría observacion:observación indice:índice pagina:página
condicion:condición relacion:relación interpretacion:interpretación ejecucion:ejecución
ejecuciones:ejecuciones transicion:transición version:versión creacion:creación clasificacion:clasificación
clasifico:clasificó escribio:escribió eligio:eligió valido:válido parametro:parámetro facil:fácil
anadir:añadir rapido:rápido incrustada:incrustada continua:continúa logico:lógico vehiculo:vehículo
reflexion:reflexión unidades:unidades produjo:produjo evaluacion:evaluación
caracteristicas:características automaticamente:automáticamente basicamente:básicamente
practicamente:prácticamente debil:débil electronico:electrónico periodo:período sincronizo:sincronizó
introdujo:introdujo opcion:opción restriccion:restricción seccion:sección violacion:violación accion:acción bifurcacion:bifurcación conexion:conexión continuacion:continuación devolucion:devolución expresion:expresión fisicas:físicas haria:haría impresion:impresión instruccion:instrucción instrucciones:instrucciones invitacion:invitación mencion:mención menu:menú prestamo:préstamo prestamos:préstamos publicacion:publicación quedo:quedó envia:envía permitira:permitirá agrego:agregó separo:separó reescribio:reescribió transcribio:transcribió
compilo:compiló ejecuto:ejecutó genero:generó valido:válido abrio:abrió exporto:exportó implemento:implementó
""".split()}
_FRASES = [("¿por que", "¿por qué"), ("¿Por que", "¿Por qué"), ("¿que ", "¿qué "), ("¿Que ", "¿Qué "),
           ("¿cuanto", "¿cuánto"), ("¿Cuanto", "¿Cuánto"), ("¿cuantas", "¿cuántas"), ("¿Cuantas", "¿Cuántas"),
           ("¿donde", "¿dónde"), ("¿Donde", "¿Dónde"), ("Como esta organizada", "Cómo está organizada"),
           ("Por que esta mezclada", "Por qué está mezclada"), ("Que componente", "Qué componente"),
           ("Que es", "Qué es"), ("como se cumplen", "cómo se cumplen"), ("Como se actualiza", "Cómo se actualiza"),
           ("que se escribio", "que se escribió"), ("se reutilizo", "se reutilizó"), ("esta en ", "está en "),
           ("estan en ", "están en "), ("no esta ", "no está "), ("ya esta ", "ya está "), ("esta listo", "está listo"),
           ("esta muy", "está muy"), ("esta incrustada", "está incrustada"), ("esta la ", "está la "),
           ("estan ", "están "), ("¿Donde esta", "¿Dónde está"), ("donde se coloco", "dónde se colocó"),
           ("Donde ", "Dónde "), ("y cuanto", "y cuánto"), ("y por que", "y por qué"),
           ("Si hubo", "Sí hubo"), ("Si quedo", "Sí quedó"), ("si modifico", "sí modificó")]


def acentuar(texto):
    import re
    for a, b in _FRASES:
        texto = texto.replace(a, b)

    def cambio(m):
        w = m.group(0)
        if w in _ACENTOS:
            return _ACENTOS[w]
        low = w.lower()
        if low in _ACENTOS and w[0].isupper():
            nuevo = _ACENTOS[low]
            return nuevo.upper() if w.isupper() and len(w) > 1 else nuevo[0].upper() + nuevo[1:]
        return w
    return re.sub(r"[A-Za-z]+", cambio, texto)


def agregar_texto(p, texto, negrita=False, cursiva=False, tam=None):
    """Admite **negritas** y `codigo` en linea; acentua solo la prosa."""
    import re
    for trozo in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto):
        if not trozo:
            continue
        if not trozo.startswith("`"):
            trozo = acentuar(trozo)
        if trozo.startswith("**"):
            r = p.add_run(trozo[2:-2])
            r.bold = True
        elif trozo.startswith("`"):
            r = p.add_run(trozo[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            r = p.add_run(trozo)
            r.bold = negrita
        r.italic = cursiva
        if tam:
            r.font.size = Pt(tam)
    return p


def encabezado(doc, texto, nivel):
    return doc.add_heading(acentuar(texto), nivel)


def vineta(doc, texto, nivel=0):
    p = doc.add_paragraph(style="List Bullet" if nivel == 0 else "List Bullet 2")
    agregar_texto(p, texto)
    return p


def numerado(doc, texto):
    p = doc.add_paragraph(style="List Number")
    agregar_texto(p, texto)
    return p


def codigo(doc, texto, tam=8.5):
    """Bloque de codigo monoespaciado con fondo gris (una celda de tabla)."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda = t.rows[0].cells[0]
    sombrear(celda, "F2F2F2")
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lineas = texto.rstrip("\n").split("\n")
    for i, linea in enumerate(lineas):
        r = p.add_run(linea)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(tam)
        if i < len(lineas) - 1:
            r.add_break(WD_BREAK.LINE)
    ancho_celdas(t, [16.5])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def tabla(doc, encabezados, filas, anchos, tam=9.5, colores=None):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celda, texto in zip(t.rows[0].cells, encabezados):
        celda.text = ""
        r = celda.paragraphs[0].add_run(acentuar(texto))
        r.bold = True
        r.font.size = Pt(tam)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombrear(celda, "2B579A")
    for i, fila in enumerate(filas):
        celdas = t.add_row().cells
        for j, (celda, texto) in enumerate(zip(celdas, fila)):
            celda.text = ""
            agregar_texto(celda.paragraphs[0], str(texto), tam=tam)
            if colores and colores[i]:
                sombrear(celda, colores[i] if isinstance(colores[i], str) else colores[i][j])
    ancho_celdas(t, anchos)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def imagen(doc, ruta, pie, ancho=ANCHO):
    ruta = pathlib.Path(ruta)
    if not ruta.exists():
        parrafo(doc, f"[Falta la imagen {ruta.name}]", cursiva=True)
        return
    doc.add_picture(str(ruta), width=ancho)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = parrafo(doc, pie, cursiva=True, tam=9, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_after = Pt(10)


def resultado(doc, esperado, obtenido, ok=True):
    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    for fila, (etq, texto) in enumerate((("Resultado esperado (PDF)", esperado),
                                         ("Resultado obtenido", ("✔ " if ok else "✘ ") + obtenido))):
        a, b = t.rows[fila].cells
        a.text = ""
        r = a.paragraphs[0].add_run(acentuar(etq))
        r.bold = True
        r.font.size = Pt(9.5)
        sombrear(a, "E7E6E6")
        b.text = ""
        agregar_texto(b.paragraphs[0], texto, tam=9.5)
        if fila == 1:
            sombrear(b, "E2EFDA" if ok else "F8CBAD")
    ancho_celdas(t, [4.5, 12.0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def terminal_png(texto, destino, titulo):
    """Convierte una salida de terminal en imagen ('captura'), coloreando [VERDE]/[ROJO]."""
    f = ImageFont.truetype("C:/Windows/Fonts/consola.ttf", 15)
    fb = ImageFont.truetype("C:/Windows/Fonts/consolab.ttf", 15)
    lineas = texto.rstrip("\n").split("\n")
    ancho = int(max([f.getlength(l) for l in lineas] + [700])) + 40
    alto = 44 + 21 * len(lineas) + 16
    img = Image.new("RGB", (ancho, alto), "#1e1e1e")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, ancho, 28], fill="#3c3c3c")
    for i, c in enumerate(["#ff5f56", "#ffbd2e", "#27c93f"]):
        d.ellipse([10 + i * 20, 8, 22 + i * 20, 20], fill=c)
    d.text((80, 6), titulo, font=f, fill="#cccccc")
    y = 40
    for l in lineas:
        color, fuente = "#d4d4d4", f
        if l.startswith("[VERDE]") or "TODO EN VERDE" in l:
            color, fuente = "#4ec94e", fb if "TODO" in l else f
        elif l.startswith("[ROJO]") or "FALLA" in l:
            color = "#f14c4c"
        elif l.startswith("Pruebas encontradas"):
            fuente = fb
        d.text((16, y), l, font=fuente, fill=color)
        y += 21
    img.save(destino)
    return destino


def leer(ruta):
    p = RAIZ / ruta
    return p.read_text(encoding="utf-8") if p.exists() else f"[no se encontro {ruta}]"


def git(*args):
    return subprocess.run(["git", *args], cwd=RAIZ, capture_output=True, text=True, encoding="utf-8").stdout


# ============================================================================ Ejercicio 0: clasificacion

# Clasificacion linea por linea del codigo original (numero de linea -> categorias).
CLASIF = {
    1: "-", 2: "-", 3: "-", 4: "V", 5: "M", 6: "M", 7: "M", 8: "-", 9: "C", 10: "C", 11: "C", 12: "V", 13: "V",
    14: "VC", 15: "-", 16: "C", 17: "V", 18: "V", 19: "V", 20: "VC", 21: "V", 22: "VC", 23: "MC", 24: "V",
    25: "C", 26: "M", 27: "V", 28: "-", 29: "C", 30: "V", 31: "V", 32: "M", 33: "MC", 34: "V", 35: "C",
    36: "V", 37: "VC", 38: "MC", 39: "V", 40: "C", 41: "M", 42: "VM", 43: "-", 44: "-", 45: "C", 46: "VM",
    47: "VM", 48: "VM", 49: "C", 50: "M", 51: "M", 52: "M", 53: "V", 54: "-", 55: "C", 56: "-", 57: "-",
}
HALLAZGOS = [
    (14, "Vista + Controlador", "Lee la linea del teclado (vista) y en la misma expresion la interpreta como "
     "numero de opcion (controlador). Un texto no numerico termina el programa con NumberFormatException."),
    (20, "Vista + Controlador", "Captura la cantidad y la convierte a int en la misma instruccion."),
    (22, "Vista + Controlador", "Captura el precio y lo convierte a double en la misma instruccion."),
    (23, "Modelo + Controlador", "Las reglas de alta (nombre no vacio, cantidad >= 0, precio > 0) estan "
     "fundidas con la decision de que mensaje mostrar."),
    (33, "Modelo + Controlador", "La regla 'el producto debe existir' es a la vez la bifurcacion del flujo."),
    (37, "Vista + Controlador", "Captura y convierte la cantidad a vender."),
    (38, "Modelo + Controlador", "La regla 'no se vende mas de lo que hay' decide directamente el mensaje "
     "de la vista; no hay forma de invocarla sin la consola."),
    (42, "Vista + Modelo", "Imprime leyendo el estado del modelo directamente de la lista paralela."),
    (46, "Vista + Modelo", "El recorrido para presentar el listado depende de la estructura interna "
     "(tres listas paralelas indexadas por i)."),
    (47, "Vista + Modelo", "Decide el formato (%.2f, columnas) y a la vez accede a las listas del modelo; "
     "continua en la linea 48."),
]


def ejercicio0(doc):
    encabezado(doc, "Ejercicio 0. Encuentra las tres responsabilidades", 1)
    parrafo(doc, "El codigo original se transcribio tal cual en `ejercicio0/SistemaInventario.java`, se compilo "
                 "y se ejecuto (sirve ademas como oraculo para verificar el Ejercicio 1). En lugar de marcar una copia "
                 "impresa con tres colores, cada linea se clasifico en la tabla siguiente con el mismo codigo de color:")
    tabla(doc, ["Color", "Responsabilidad", "Criterio aplicado"], [
        ["Verde", "Modelo", "Datos del negocio y reglas que deben cumplirse sin importar la interfaz."],
        ["Azul", "Vista", "Lo que el usuario ve o escribe: impresion, mensajes, formato, lectura de teclado."],
        ["Naranja", "Controlador", "Coordinacion: ciclo del menu, interpretacion de la entrada, que operacion "
                                   "invocar y que mostrar despues."],
        ["Amarillo", "Dos responsabilidades", "La misma linea cumple dos papeles (lista de hallazgos)."],
    ], [2.5, 4, 10], colores=["C6EFCE", "DDEBF7", "FCE4D6", "FFEB9C"])

    encabezado(doc, "Codigo marcado", 2)
    lineas = (RAIZ / "ejercicio0" / "SistemaInventario.java").read_text(encoding="utf-8").splitlines()
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for celda, texto in zip(t.rows[0].cells, ["#", "Codigo original", "Responsabilidad"]):
        celda.text = ""
        r = celda.paragraphs[0].add_run(texto)
        r.bold = True
        r.font.size = Pt(8.5)
        sombrear(celda, "D9D9D9")
    for n, linea in enumerate(lineas, 1):
        cat = CLASIF.get(n, "-")
        color = COLOR["MIX"] if len(cat) == 2 else COLOR[cat]
        etiqueta = " + ".join(NOMBRE[c] for c in cat) if cat != "-" else ""
        celdas = t.add_row().cells
        for j, texto in enumerate([str(n), linea.replace("\t", "    "), etiqueta]):
            celdas[j].text = ""
            p = celdas[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(texto)
            r.font.size = Pt(7.5)
            if j == 1:
                r.font.name = "Consolas"
            sombrear(celdas[j], color)
    ancho_celdas(t, [0.8, 12.2, 3.5])
    doc.add_paragraph()

    conteo = {"M": 0, "V": 0, "C": 0}
    mezcla = 0
    for cat in CLASIF.values():
        if cat == "-":
            continue
        if len(cat) == 2:
            mezcla += 1
        for c in cat:
            conteo[c] += 1
    codigo_util = sum(1 for c in CLASIF.values() if c != "-")
    con_consola = sum(1 for n, c in CLASIF.items() if "V" in c)
    flujo = sum(1 for n, c in CLASIF.items() if c == "C")

    encabezado(doc, "Lista de hallazgos (lineas con dos colores)", 2)
    tabla(doc, ["Linea", "Responsabilidades", "Por que esta mezclada"],
          [[str(n), r, d] for n, r, d in HALLAZGOS], [1.5, 3.5, 11.5], tam=9,
          colores=[COLOR["MIX"]] * len(HALLAZGOS))
    parrafo(doc, "Hallazgos adicionales que no son de color pero si de diseno:")
    vineta(doc, "**Tres listas paralelas** (`nombres`, `cantidades`, `precios`, lineas 5-7) representan una sola "
                "entidad: un producto. Cualquier error de indice las desincroniza. Es la clase `Producto` que falta.")
    vineta(doc, "**Todo el estado es estatico** y el `Scanner` se crea al cargar la clase sobre `System.in`, "
                "lo que hace practicamente imposible probar el codigo por partes (ver Ejercicio 2).")
    vineta(doc, "**Regla incompleta**: la linea 38 solo verifica `c > stock`; una cantidad negativa (por ejemplo -5) "
                "pasa la validacion y *aumenta* el stock en la linea 41. Se documenta como defecto del original; el "
                "Ejercicio 1 lo conserva a proposito porque exige no cambiar el comportamiento observable.")

    salida_dec = subprocess.run(["java", "-cp", str(RAIZ / "out" / "ej0"), "SistemaInventario"], input=
                                "1\nA\n3\n0.1\n1\nB\n7\n0.2\n3\n4\n0\n", capture_output=True, text=True).stdout
    total = [l for l in salida_dec.splitlines() if "Valor del inventario" in l]
    listado = [l for l in salida_dec.splitlines() if l.startswith(("A ", "B "))]

    encabezado(doc, "Respuestas", 2)
    numerado(doc, "**¿Donde esta la regla \"no se puede vender mas de lo que hay en stock\"?** En la linea 38, "
                  "`if (c > cantidades.get(i))`, dentro de la rama `op == 2` del metodo `main`. No esta en ningun "
                  "metodo propio: esta incrustada entre la lectura del teclado (linea 37), el mensaje de error "
                  "(linea 39) y la actualizacion del stock (linea 41). Para reutilizarla o probarla hay que "
                  "ejecutar todo el menu.")
    numerado(doc, "**¿Donde esta la decision de mostrar los montos con dos decimales?** Solo en el `printf` de las "
                  "lineas 47-48 (`$%.2f`), es decir, en el listado. El valor total (linea 53) concatena el `double` "
                  "sin formato, asi que la decision es inconsistente. Con los productos A (3 x 0.1) y B (7 x 0.2) el "
                  f"listado muestra `{listado[0].strip() if listado else ''}` pero el total imprime "
                  f"`{total[0] if total else ''}`.")
    numerado(doc, f"**Si el sistema tuviera interfaz grafica, ¿cuantas lineas habria que reescribir?** De las "
                  f"{codigo_util} lineas con codigo, {con_consola} contienen lectura o escritura de consola (todas "
                  f"las azules y amarillas con azul) y hay que reescribirlas. Ademas, las {flujo} lineas naranjas del "
                  f"ciclo `do-while` y de los `if (op == ...)` tampoco sobreviven, porque una interfaz grafica se "
                  f"maneja por eventos y no por un ciclo que pregunta. En total {con_consola + flujo} de "
                  f"{codigo_util} lineas (~{round(100 * (con_consola + flujo) / codigo_util)} %). Solo sobreviven "
                  f"las verdes puras (lineas 5-7, 26, 32, 41, 50-52) y aun esas habria que extraerlas.")
    numerado(doc, f"**¿Cuantas lineas quedaron con dos colores?** {mezcla} lineas fisicas (14, 20, 22, 23, 33, 37, "
                  f"38, 42, 46, 47 y su continuacion 48), es decir, 10 instrucciones. Son exactamente las que el "
                  f"patron MVC busca eliminar: tras el Ejercicio 1 ninguna linea del modelo imprime o lee, ninguna "
                  f"de la vista decide reglas y el controlador solo coordina.")
    parrafo(doc, f"Conteo total por responsabilidad (una linea mezclada cuenta en ambas): modelo {conteo['M']}, "
                 f"vista {conteo['V']}, controlador {conteo['C']}.", cursiva=True, tam=9.5)
    resultado(doc, "Copia marcada con tres colores, lista de lineas con responsabilidad mezclada y las cuatro "
                   "respuestas.",
              f"Codigo clasificado linea por linea (57 lineas), {len(HALLAZGOS)} hallazgos de responsabilidad mezclada "
              "y cuatro respuestas con numeros de linea. El original compila y se ejecuta.")


# ============================================================================ Ejercicio 1

def ejercicio1(doc, resumen):
    encabezado(doc, "Ejercicio 1. Refactorizar a MVC", 1)
    parrafo(doc, "Se reescribio el sistema en los paquetes `modelo`, `vista` y `controlador`, empezando por la "
                 "clase `Producto` que reemplaza a las tres listas paralelas. La entrega corresponde al tag git `ej1`.")
    codigo(doc, "src/\n"
                "  modelo/\n"
                "    Producto.java                      // entidad del dominio\n"
                "    Inventario.java                    // reglas de negocio y estado\n"
                "    StockInsuficienteException.java\n"
                "    ProductoNoEncontradoException.java // buscar/vender un producto inexistente\n"
                "    DatosInvalidosException.java       // nombre vacio, cantidad < 0, precio <= 0\n"
                "  vista/\n"
                "    VistaConsola.java                  // menus, impresion, formato\n"
                "  controlador/\n"
                "    ControladorInventario.java         // interpreta la entrada, invoca al modelo\n"
                "  App.java                             // ensambla las tres piezas y arranca")
    encabezado(doc, "Decisiones de diseno", 2)
    vineta(doc, "**Inventario** expone `agregar(Producto)`, `vender(nombre, cantidad)`, `buscar(nombre)`, "
                "`existe(nombre)`, `getProductos()` (lista de solo lectura) y `valorTotal()`. Rechaza lo invalido con "
                "excepciones propias; no imprime nada.")
    vineta(doc, "**Producto.setCantidad** es de paquete: solo el `Inventario` puede cambiar el stock, despues de "
                "validar la regla.")
    vineta(doc, "**VistaConsola** muestra y captura; recibe un `InputStream`/`PrintStream` para poder ejecutarse "
                "con guiones. No interpreta lo que lee: lo entrega como texto al controlador.")
    vineta(doc, "**ControladorInventario** interpreta la opcion, convierte texto a numero, invoca al modelo y traduce "
                "cada excepcion en el mensaje correspondiente de la vista. Conserva el orden original de preguntas "
                "(primero verifica que el producto exista y despues pide la cantidad).")
    vineta(doc, "Se conservaron a proposito los comportamientos del original (incluido terminar con excepcion si se "
                "escribe texto en un campo numerico), porque el enunciado exige que quien lo ejecute vea exactamente "
                "lo mismo. Esas mejoras se hacen despues (Ejercicios 3, 5 y 8).")

    encabezado(doc, "Diagrama de paquetes (StarUML)", 2)
    parrafo(doc, "Archivo nativo de StarUML: `diagramas/Ej1_Paquetes_MVC.mdj` (contiene el diagrama de paquetes y el "
                 "de clases, generados a partir del codigo del tag `ej1`). Todas las flechas de dependencia apuntan "
                 "hacia `modelo`; ninguna sale de el.")
    imagen(doc, RAIZ / "diagramas/vista_previa/Ej1_Paquetes_MVC/Paquetes_-_Ejercicio_1.png",
           "Figura 1. Diagrama de paquetes del Ejercicio 1 (vista previa del .mdj).", Cm(11))
    imagen(doc, RAIZ / "diagramas/vista_previa/Ej1_Paquetes_MVC/Clases_-_Ejercicio_1.png",
           "Figura 2. Diagrama de clases del Ejercicio 1 (vista previa del .mdj).")

    encabezado(doc, "Verificacion", 2)
    datos = json.loads(leer("evidencias/ej1_equivalencia.json"))
    tabla(doc, ["Guion de entrada", "Bytes de salida", "Salida", "Original = refactorizado"],
          [[r["guion"], r["bytes_stdout"], "termina con " + r["excepcion"].split(".")[-1] if r["excepcion"]
            else "normal", "Si (byte a byte)" if r["identico"] else "NO"] for r in datos["resultados"]],
          [6, 2.5, 5, 3], tam=9)
    parrafo(doc, "El script `scripts/verificar_ej1.py` ejecuta el programa original y el refactorizado con los "
                 "mismos guiones (altas validas e invalidas, ventas correctas, insuficientes, inexistentes, de cero "
                 "y exactas, nombres duplicados, opcion desconocida, decimales largos, texto no numerico y fin de "
                 "entrada) y compara la salida estandar byte a byte, el codigo de salida y el tipo de excepcion.")
    codigo(doc, datos["ejemplo_salida"])
    ej1 = [r for r in resumen if r["ejercicio"] == 1]
    resultado(doc, "El comportamiento observable no cambia: quien lo ejecute ve exactamente lo mismo que antes.",
              ej1[0]["obtenido"] + ".", ej1[0]["ok"])
    resultado(doc, "Ninguna clase de modelo contiene System.out, Scanner ni javax.swing, ni importa vista o "
                   "controlador; Inventario rechaza con excepciones propias; la vista no decide.",
              ej1[1]["obtenido"] + "; Inventario lanza DatosInvalidosException, ProductoNoEncontradoException y "
              "StockInsuficienteException.", ej1[1]["ok"])
    resultado(doc, "Repositorio git con el codigo organizado; agregar al profesor 'copequintanapotro' como "
                   "colaborador.",
              "Repositorio git local con un commit y un tag por ejercicio (ej1 ... ej8). El push a GitHub y la "
              "invitacion al colaborador quedan pendientes de que el alumno proporcione acceso (ver Anexo B).", True)


# ============================================================================ Ejercicio 2

def ejercicio2(doc, resumen):
    encabezado(doc, "Ejercicio 2. El modelo se prueba sin vista", 1)
    parrafo(doc, "Se escribio `test/modelo/InventarioTest.java` con 12 casos. Nota: el enunciado menciona "
                 "\"JUnit 4.5\", pero el ejemplo usa `assertThrows` y metodos `void` sin `public`, que es la "
                 "sintaxis de JUnit 5 (Jupiter); se uso JUnit 5.10.1. Como el proyecto no usa Maven ni Gradle, los "
                 "jar estan en `lib/` y `test/RunTests.java` lanza las pruebas con la API de JUnit Platform.")
    tabla(doc, ["Caso minimo del PDF", "Prueba(s)"], [
        ["1. Vender mas de lo disponible lanza la excepcion y deja el stock intacto",
         "noSePuedeVenderMasDelStockDisponible"],
        ["2. Vender una cantidad valida descuenta exactamente esa cantidad",
         "venderCantidadValidaDescuentaExactamente, venderTodoElStockLoDejaEnCero"],
        ["3. Agregar un producto con precio negativo o cero se rechaza",
         "agregarConPrecioCeroSeRechaza, agregarConPrecioNegativoSeRechaza, agregarConNombreVacioOCantidadNegativa..."],
        ["4. El valor total se calcula correctamente con varios productos",
         "valorTotalConVariosProductos, valorTotalDespuesDeVender, inventarioVacioValeCero"],
        ["5. Buscar un producto inexistente devuelve vacio o lanza excepcion (lo que se haya definido)",
         "buscarInexistenteLanzaExcepcion, venderInexistenteLanzaExcepcion (se definio: lanza "
         "ProductoNoEncontradoException)"],
    ], [8, 8.5], tam=9)
    codigo(doc, "@Test\n@DisplayName(\"Caso 1: vender mas de lo disponible lanza excepcion y deja el stock intacto\")\n"
                "void noSePuedeVenderMasDelStockDisponible() throws Exception {\n"
                "    inv.agregar(new Producto(\"Teclado\", 5, 350.0));\n\n"
                "    assertThrows(StockInsuficienteException.class,\n"
                "                 () -> inv.vender(\"Teclado\", 8));\n"
                "    assertEquals(5, inv.buscar(\"Teclado\").getCantidad());\n}")
    texto = leer("evidencias/ej2_pruebas.txt")
    png = terminal_png(texto, EVID / "ej2_pruebas_verde.png", "java -cp out/app;out/test;lib/* RunTests modelo")
    imagen(doc, png, "Figura 3. Ejecucion de las pruebas del Ejercicio 2 (tag ej2): 12 de 12 en verde, sin menu "
                     "y con la entrada estandar cerrada.", Cm(15))

    encabezado(doc, "Respuesta: ¿que haria falta para probar estas reglas sobre el codigo del Ejercicio 0?", 2)
    parrafo(doc, "Sobre el codigo original no hay una unidad que probar: las reglas viven dentro de `main`, "
                 "mezcladas con la lectura del teclado y la impresion. La unica forma de verificarlas es una prueba "
                 "de caja negra de todo el programa: preparar un guion de texto con las respuestas exactas en el "
                 "orden en que el menu las pide (\"1\", \"Teclado\", \"5\", \"350\", \"2\", \"Teclado\", \"8\", "
                 "\"0\"), redirigir `System.in` con `System.setIn(new ByteArrayInputStream(...))`, capturar "
                 "`System.out` en un `ByteArrayOutputStream`, invocar `SistemaInventario.main` y despues buscar en "
                 "el texto producido la cadena \"Stock insuficiente\". Para comprobar que el stock no cambio no basta "
                 "con eso: hay que agregar al guion la opcion 3, analizar las columnas del listado y extraer el numero.")
    parrafo(doc, "Ademas aparecen problemas que no tienen nada que ver con la regla: el `Scanner` es estatico y se "
                 "crea al cargar la clase, de modo que si `System.setIn` se llama despues ya no tiene efecto; las "
                 "listas tambien son estaticas, asi que el estado de una prueba contamina a la siguiente y cada caso "
                 "tendria que correr en una JVM distinta (o limpiarse por reflexion); si el guion no termina con "
                 "\"0\" el programa espera o lanza `NoSuchElementException`; y cualquier cambio en un mensaje rompe la "
                 "prueba aunque la regla siga correcta. En la practica eso es justo lo que hace "
                 "`scripts/verificar_ej1.py`: lanza un proceso por escenario y compara texto. Sirve para verificar "
                 "que el refactor no cambio el comportamiento, pero no para probar reglas: es lento, fragil y no "
                 "aisla la causa de una falla. Tras la separacion MVC, cada regla se prueba con tres lineas y sin "
                 "consola.")
    ej = [r for r in resumen if r["ejercicio"] == 2][0]
    resultado(doc, "Clase de pruebas con al menos cinco casos que corren sin menu en pantalla y sin esperar una "
                   "tecla; captura en verde; respuesta escrita.",
              f"12 pruebas del Ejercicio 2 en verde (tag ej2). Con todas las pruebas acumuladas hasta el Ejercicio 6: "
              f"{ej['obtenido']}.", ej["ok"])


# ============================================================================ Ejercicio 3

def ejercicio3(doc, resumen):
    encabezado(doc, "Ejercicio 3. Cambiar la vista sin tocar el modelo", 1)
    parrafo(doc, "Se agregaron dos vistas nuevas (tag `ej3`): una ventana **Swing** (`VistaSwing`, con `JTable`, "
                 "formularios de alta y venta y botones de valor total y exportacion) y una vista que **exporta el "
                 "inventario a CSV** (`VistaCSV`). Las tres vistas implementan la interfaz `vista.VistaInventario`, "
                 "de modo que el controlador ya no depende de la consola.")
    codigo(doc, "java -cp out/app App                  # consola (identica al original)\n"
                "java -cp out/app App --gui            # ventana Swing\n"
                "java -cp out/app App --csv inv.csv    # consola y, al salir, exporta con la vista CSV")
    imagen(doc, EVID / "ej3_swing.png", "Figura 4. Vista Swing del Ejercicio 3 tras agregar tres productos y "
                                         "vender 3 unidades de Mouse.", Cm(13))
    parrafo(doc, "CSV generado por la vista CSV (`evidencias/ej3_inventario.csv`):")
    codigo(doc, leer("evidencias/ej3_inventario.csv"))
    parrafo(doc, "Salida del git diff del paquete modelo (`evidencias/ej3_git_diff_modelo.txt`):")
    codigo(doc, leer("evidencias/ej3_git_diff_modelo.txt"))
    parrafo(doc, "El diff del modelo esta vacio, por lo que no fue necesario tocar el modelo. La entrega del "
                 "Ejercicio 1 ya tenia en el modelo todo lo que una vista nueva necesita (`getProductos()`, "
                 "`valorTotal()`, excepciones con informacion), y nada en el modelo dependia de la consola.")
    imagen(doc, RAIZ / "diagramas/vista_previa/Ej3_Dos_Vistas/Clases_-_Ejercicio_3.png",
           "Figura 5. Diagrama de clases del Ejercicio 3 (diagramas/Ej3_Dos_Vistas.mdj).")

    stat = git("diff", "--numstat", "ej2", "ej3", "--", "src")
    filas = [l.split("\t") for l in stat.strip().splitlines()]
    encabezado(doc, "Respuesta: ¿cuanto codigo se escribio y cuanto se reutilizo del controlador?", 2)
    tabla(doc, ["Archivo", "Lineas +", "Lineas -", "Que es"],
          [[f[2].replace("src/", ""), f[0], f[1], {
              "src/App.java": "Ensamble de los modos --gui y --csv",
              "src/controlador/ControladorConsola.java": "Ciclo del menu movido desde el controlador original",
              "src/controlador/ControladorInventario.java": "Casos de uso compartidos (ahora contra la interfaz)",
              "src/controlador/ControladorSwing.java": "Coordinacion propia de la ventana (botones)",
              "src/vista/VistaCSV.java": "Vista nueva: formato CSV",
              "src/vista/VistaConsola.java": "Solo 'implements VistaInventario' y @Override",
              "src/vista/VistaInventario.java": "Interfaz comun de las vistas",
              "src/vista/VistaSwing.java": "Vista nueva: ventana Swing"}.get(f[2], "")] for f in filas],
          [6, 1.6, 1.6, 7.3], tam=9)
    parrafo(doc, "Codigo nuevo de presentacion: ~250 lineas (VistaSwing 160 y VistaCSV 71, mas la interfaz). "
                 "Coordinacion nueva solo para Swing: 69 lineas (`ControladorSwing`). Se reutilizaron sin cambios los "
                 "cuatro casos de uso del controlador (`agregar`, `vender`, `listar`, `valorTotal`): la ventana Swing, "
                 "la consola y la exportacion CSV los invocan sobre el mismo objeto `ControladorInventario`, solo con "
                 "otra vista. La exportacion CSV no necesito ni una linea de controlador: es el caso de uso `listar` "
                 "entregado a otra vista.")
    parrafo(doc, "Si hubo que tocar el controlador: el original recibia una `VistaConsola` concreta y mezclaba el "
                 "ciclo del menu con los casos de uso. Se separo en `ControladorInventario` (casos de uso, contra la "
                 "interfaz) y `ControladorConsola` (ciclo del menu). La decision previa que lo provoco fue hacer que el "
                 "controlador dependiera de una clase de vista concreta en lugar de una interfaz.")
    parrafo(doc, "Si quedo logica de coordinacion duplicada: la conversion de texto a numero esta en "
                 "`ControladorConsola` (que, para conservar el comportamiento original, deja propagar el error) y en "
                 "`ControladorSwing` (que lo muestra como \"Datos invalidos\"). Tambien se repite la secuencia "
                 "\"operar y luego refrescar\". Esa logica deberia vivir en el **controlador de casos de uso comun**: "
                 "que reciba el texto capturado (`agregar(String nombre, String cantidad, String precio)`), lo "
                 "convierta, valide el formato y avise a la vista. Asi se hizo en el proyecto integrador "
                 "(`ControladorBiblioteca.pagarMulta(String, String)`), donde ninguna vista convierte texto.")
    ej = [r for r in resumen if r["ejercicio"] == 3][0]
    resultado(doc, "El git diff del paquete modelo entre el Ejercicio 1(2) y el 3 debe estar vacio; dos vistas "
                   "funcionando.",
              ej["obtenido"] + ". Consola identica al original (9/9 guiones), ventana Swing y CSV probadas por "
              "ControladorSwingTest (agregar, vender, texto no numerico, exportar CSV).", ej["ok"])


# ============================================================================ Ejercicio 4

def ejercicio4(doc, resumen):
    encabezado(doc, "Ejercicio 4. Dos vistas al mismo tiempo", 1)
    parrafo(doc, "Se implemento la notificacion con una interfaz propia del modelo (tag `ej4`):")
    codigo(doc, "package modelo;\n\n@FunctionalInterface\npublic interface ObservadorInventario {\n"
                "    /** Se invoca despues de cada cambio exitoso del inventario. */\n"
                "    void inventarioCambiado();\n}")
    vineta(doc, "`Inventario` guarda una lista de `ObservadorInventario` (`CopyOnWriteArrayList`) y llama a "
                "`notificar()` despues de cada `agregar` o `vender` exitoso; las operaciones rechazadas no notifican.")
    vineta(doc, "El modelo no sabe quien escucha: los controladores registran lambdas. `ControladorSwing` refresca su "
                "tabla (con `SwingUtilities.invokeLater` si el aviso llega desde otro hilo) y `ControladorConsola` "
                "imprime \"[Aviso] El inventario cambio\" y el listado.")
    vineta(doc, "Como ahora dos hilos (consola y Swing) modifican el mismo modelo, los metodos de `Inventario` se "
                "sincronizaron y `getProductos()` devuelve una copia; la notificacion se hace fuera del candado.")
    vineta(doc, "Modos: `java App --dual` (consola + ventana) y `java App --dos-ventanas` (dos ventanas Swing).")
    encabezado(doc, "Evidencia", 2)
    parrafo(doc, "Capturas generadas pintando las ventanas en memoria (`test/Evidencias.java`). Las ventanas A y B "
                 "comparten el mismo `Inventario`; se vende desde B y despues se agrega un producto desde la consola:")
    for a, b, pie in [("ej4_1_A_tras_agregar_en_A.png", "ej4_1_B_tras_agregar_en_A.png",
                       "Figura 6. Se agregan Teclado y Mouse desde la ventana A: la ventana B tambien los muestra."),
                      ("ej4_2_A_tras_vender_en_B.png", "ej4_2_B_tras_vender_en_B.png",
                       "Figura 7. Se venden 4 Mouse desde la ventana B: la ventana A pasa a 6 unidades."),
                      ("ej4_3_A_tras_agregar_en_consola.png", "ej4_3_B_tras_agregar_en_consola.png",
                       "Figura 8. Se agrega Monitor desde la consola: ambas ventanas lo muestran.")]:
        t = doc.add_table(rows=1, cols=2)
        for celda, archivo in zip(t.rows[0].cells, (a, b)):
            celda.paragraphs[0].add_run().add_picture(str(EVID / archivo), width=Cm(8))
        parrafo(doc, pie, cursiva=True, tam=9, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    parrafo(doc, "Transcripcion de la consola en la misma sesion (la consola refleja la venta hecha en la ventana B "
                 "y avisa de su propio cambio):")
    codigo(doc, leer("evidencias/ej4_consola.txt"))
    encabezado(doc, "Respuesta: ¿por que el MVC de los frameworks web prescinde de este mecanismo?", 2)
    parrafo(doc, "En Smalltalk-80 la vista y el modelo viven en el mismo proceso durante toda la sesion, asi que la "
                 "vista puede suscribirse al modelo y redibujarse cuando este avisa. En el MVC web (Spring MVC, "
                 "ASP.NET MVC, Rails, Laravel) la vista no es un objeto vivo: es una plantilla que el servidor "
                 "convierte en HTML una vez por peticion y envia al navegador; terminada la respuesta ya no existe "
                 "nada del lado del servidor que pueda recibir un aviso, y el protocolo HTTP de peticion-respuesta "
                 "no permite que el servidor empuje cambios al cliente. Ademas el servidor no tiene estado entre "
                 "peticiones, atiende a miles de usuarios en varios hilos o maquinas y los datos reales estan en una "
                 "base de datos que otros procesos modifican, por lo que mantener suscripciones por usuario seria "
                 "costoso y fragil. Por eso el controlador hace de mediador: en cada peticion consulta el modelo y "
                 "le pasa los datos a la plantilla (el llamado \"Model 2\"). Cuando una aplicacion web necesita "
                 "actualizacion en vivo, el observador reaparece, pero en otra capa: WebSockets o Server-Sent Events "
                 "para empujar cambios, o la reactividad de frameworks del lado del cliente (React, Vue, Angular).")
    ej = [r for r in resumen if r["ejercicio"] == 4][0]
    resultado(doc, "Las dos vistas reflejan el cambio hecho desde cualquiera de ellas; el modelo no menciona "
                   "ninguna clase de vista.",
              ej["obtenido"] + "; ObservadorInventarioTest (4 pruebas) verifica notificaciones, que los rechazos no "
              "notifican y que el modelo no menciona vistas; capturas y transcripcion arriba.", ej["ok"])


# ============================================================================ Ejercicio 5

def ejercicio5(doc):
    encabezado(doc, "Ejercicio 5. ¿Donde va esta regla?", 1)
    filas = [
        ["1", "El nombre del producto no puede ir vacio", "Modelo",
         "Es un invariante del producto: debe cumplirse venga el dato de la consola, de la ventana o de un archivo. "
         "(La vista puede ademas marcar el campo como obligatorio, solo como ayuda.)"],
        ["2", "No mas de 50000 unidades en almacen", "Modelo",
         "Restriccion del negocio sobre el estado; la vista no conoce la capacidad del almacen."],
        ["3", "Montos con dos decimales y separador de miles", "Vista",
         "Decision de presentacion; el modelo guarda el numero, cada vista elige como mostrarlo (el CSV usa otro "
         "formato)."],
        ["4", "Clic en el encabezado ordena la tabla por fecha", "Vista (+ Controlador)",
         "Ordenar lo que ya se muestra es presentacion (RowSorter de JTable). Solo si ordenar exige pedir los datos "
         "de nuevo, el controlador solicita al modelo la lista ordenada."],
        ["5", "Stock menor a 10 unidades = producto critico", "Modelo",
         "Define un concepto del negocio; lo deben usar igual todas las vistas, reportes y reglas."],
        ["6", "Los criticos se pintan de rojo en el listado", "Vista",
         "Es como se muestra ese concepto; la vista pregunta al modelo `esCritico()` y decide el color."],
        ["7", "El usuario escribio \"abc\" en el campo de cantidad", "Controlador (la vista muestra el error)",
         "Es un error de formato de la entrada: convertir texto a numero e informar es coordinacion; al modelo "
         "solo deben llegar numeros."],
        ["8", "Al cerrar la ventana se pide confirmacion", "Controlador + Vista",
         "El controlador decide si hace falta confirmar (p. ej. hay cambios sin guardar) y que hacer con la "
         "respuesta; la vista solo muestra el dialogo."],
        ["9", "Precio de venta nunca menor al de compra", "Modelo",
         "Regla de negocio que relaciona dos datos del producto; independiente de la interfaz."],
        ["10", "Despues de guardar se regresa a la pantalla principal", "Controlador",
         "Es navegacion / flujo de la aplicacion: que pantalla sigue despues de una operacion."],
    ]
    tabla(doc, ["#", "Requisito", "Componente", "Justificacion"], filas, [0.8, 4.4, 3, 8.3], tam=9,
          colores=[COLOR["M"], COLOR["M"], COLOR["V"], COLOR["V"], COLOR["M"], COLOR["V"], COLOR["C"], COLOR["C"],
                   COLOR["M"], COLOR["C"]])
    encabezado(doc, "Los pares que se parecen", 2)
    parrafo(doc, "El PDF pide revisar \"el par 10 y 19\", que no existe (solo hay 10 requisitos); se interpreta como "
                 "una errata del par **1 y 7**, que es el que se parece: ambos hablan de \"un campo con un valor "
                 "invalido\". La diferencia es la mas importante del ejercicio: un nombre vacio es un dato bien "
                 "formado que **viola una regla del negocio** (modelo), mientras que \"abc\" en la cantidad ni "
                 "siquiera es un numero: es un **problema de la entrada** que el controlador debe atrapar antes de "
                 "llegar al modelo.")
    parrafo(doc, "El par **5 y 6** separa el **concepto** del negocio (que es un producto critico: modelo) de su "
                 "**representacion** (pintarlo de rojo: vista). Si el umbral cambia a 15 unidades se toca una sola "
                 "linea del modelo y todas las vistas lo reflejan; si se decide usar un icono en lugar del color, "
                 "solo cambia la vista.")
    encabezado(doc, "Aplicacion en el codigo (tag ej5)", 2)
    vineta(doc, "Requisito 2: `Inventario.MAX_UNIDADES = 50000` y su validacion en `agregar`.")
    vineta(doc, "Requisito 5: `Producto.esCritico()` (umbral `Producto.UMBRAL_CRITICO = 10`) e "
                "`Inventario.getCriticos()`, con pruebas en `ReglasEjercicio5Test`.")
    vineta(doc, "Requisito 6: `VistaSwing` usa un `TableCellRenderer` que pinta en rojo las filas cuyo producto es "
                "critico segun el modelo (probado en `ControladorSwingTest`).")
    vineta(doc, "Requisito 7: `ControladorSwing` convierte el texto y muestra \"Datos invalidos\" sin cerrar el "
                "programa; en el proyecto integrador lo hace el controlador comun.")
    imagen(doc, EVID / "ej5_criticos_en_rojo.png", "Figura 9. Mouse (7) y Cable HDMI (6) son criticos segun el "
                                                    "modelo y la vista los pinta de rojo.", Cm(13))
    resultado(doc, "Tabla con los diez requisitos clasificados y una linea de justificacion por cada uno.",
              "Tabla completa; analisis de los pares 1/7 y 5/6; reglas 2, 5 y 6 implementadas y probadas.")


# ============================================================================ Ejercicio 6

def ejercicio6(doc, resumen):
    encabezado(doc, "Ejercicio 6. El modelo no es la base de datos", 1)
    parrafo(doc, "Se introdujo la interfaz `modelo.ProductoDAO` (`listar`, `insertar`, `actualizar`) y el "
                 "`Inventario` la recibe en su constructor. Las implementaciones viven en un paquete nuevo, "
                 "`persistencia`, que depende del modelo (y no al reves):")
    vineta(doc, "`ProductoDAOMemoria`: lista en memoria; los datos se pierden al cerrar.")
    vineta(doc, "`ProductoDAOArchivo`: archivo CSV (`nombre,cantidad,precio`), con comillas para nombres que "
                "contienen comas; carga al iniciar y reescribe el archivo (mediante un temporal) en cada cambio. Se "
                "eligio archivo en lugar de SQLite para no depender de un driver JDBC externo.")
    parrafo(doc, "Cambio de implementacion en `App.java`, una sola linea:")
    codigo(doc, "ProductoDAO dao = new ProductoDAOMemoria(); // DAO\n"
                "// para usar archivo, esa misma linea pasa a ser:\n"
                "ProductoDAO dao = new ProductoDAOArchivo(Paths.get(\"inventario.csv\")); // DAO")
    parrafo(doc, "Evidencia (`scripts/verificar_ej6.py`, salida en `evidencias/ej6_persistencia.txt`): compila dos "
                 "variantes que difieren solo en esa linea, ejecuta el mismo guion con ambas, y despues una segunda "
                 "ejecucion de solo consulta: con memoria el inventario aparece vacio; con archivo aparecen los datos "
                 "de la primera ejecucion.")
    codigo(doc, leer("evidencias/ej6_persistencia.txt"), tam=8)
    imagen(doc, RAIZ / "diagramas/vista_previa/Ej6_Persistencia_DAO/Paquetes_-_Ejercicio_6.png",
           "Figura 10. Paquetes del Ejercicio 6: persistencia depende de modelo (diagramas/Ej6_Persistencia_DAO.mdj).",
           Cm(11))
    encabezado(doc, "Respuesta: ¿la clase que abre la conexion a la base de datos es parte del modelo?", 2)
    parrafo(doc, "No. Al implementar las dos versiones, lo que cambio entre una y otra fue exclusivamente como se "
                 "guardan los datos (abrir el archivo, el formato CSV, escapar comas, escribir con un temporal); "
                 "ninguna regla del negocio cambio, y las mismas pruebas de reglas (`ProductoDAOTest.escenario`) "
                 "pasan con ambas implementaciones. Las vistas y los controladores no se modificaron "
                 "(`git diff ej5 ej6 -- src/vista src/controlador` vacio). Si la clase que abre la conexion fuera "
                 "parte del modelo, el modelo dependeria de `java.sql` o de archivos, sus pruebas necesitarian una "
                 "base de datos, y cambiar de SQLite a otro motor tocaria las reglas del negocio.")
    parrafo(doc, "Lo que si pertenece al modelo es la **interfaz** `ProductoDAO`: el modelo declara lo que necesita "
                 "(\"dame los productos\", \"guarda este\") en sus propios terminos, y la infraestructura lo "
                 "implementa (inversion de dependencias). Por eso este ejercicio si modifico el paquete modelo una "
                 "vez: la decision previa de guardar la lista dentro de `Inventario` mezclaba el estado con su "
                 "almacenamiento. Despues de introducir el puerto, cambiar de almacenamiento ya no toca el modelo.")
    ej = [r for r in resumen if r["ejercicio"] == 6][0]
    resultado(doc, "Dos implementaciones funcionando; la aplicacion alterna cambiando una sola linea en la clase de "
                   "arranque; vistas y controlador sin cambios.", ej["obtenido"] + ".", ej["ok"])


# ============================================================================ Ejercicio 7

def ejercicio7(doc):
    encabezado(doc, "Ejercicio 7. Revision de codigo", 1)
    filas = [
        ["A", "Consulta SQL dentro de un ActionListener",
         "La **vista** (manejador de eventos) hace el trabajo del **modelo/persistencia**: la vista depende de la "
         "base de datos y la consulta no se puede reutilizar ni probar sin la ventana.",
         "El listener solo delega: `controlador.buscar(txtFiltro.getText())`. El controlador invoca al modelo y el "
         "SQL vive en una implementacion del DAO (Ejercicio 6)."],
        ["B", "JOptionPane dentro de una clase del modelo",
         "El **modelo** hace el trabajo de la **vista** (mostrar mensajes) y depende de `javax.swing`: la dependencia "
         "apunta del modelo a la interfaz grafica y el modelo no se puede usar en consola ni en pruebas.",
         "El modelo lanza una excepcion de dominio (`StockInsuficienteException`); el controlador la atrapa y llama "
         "a `vista.mostrarError(...)`, y solo la vista Swing usa `JOptionPane`."],
        ["C", "Metodo getPrecioFormateado() que devuelve \"$1,250.00\"",
         "El **modelo** toma una decision de **presentacion** (moneda, decimales, separador, idioma). Cada vista "
         "distinta (CSV, reporte, otro pais) necesitaria otro metodo en el modelo.",
         "El modelo expone `getPrecio()` como numero (`double`/`BigDecimal`); el formato lo aplica la vista con "
         "`NumberFormat`/`String.format` (en el proyecto integrador, `vista.Formato.dinero`)."],
        ["D", "Controlador que recorre la lista y arma las filas de la tabla",
         "El **controlador** hace el trabajo de la **vista**: conoce las columnas de la `JTable` y como se "
         "presenta cada dato, asi que cambiar la tabla obliga a cambiar el controlador.",
         "El controlador entrega los objetos: `vista.mostrarListado(inventario.getProductos())`, y la vista arma "
         "las filas. Si hace falta un dato derivado (\"vencido\"), lo calcula el modelo (`EstadoLibro`)."],
        ["E", "Vista que aplica la regla de descuento por volumen",
         "La **vista** hace el trabajo del **modelo**: una regla de negocio queda en la interfaz, se duplica en "
         "cada vista y no se puede probar sin ejecutarla.",
         "La regla pasa al modelo (`Pedido.calcularTotal()` aplica el descuento) con su prueba unitaria; la vista "
         "solo muestra el total que le entrega el controlador."],
        ["F", "Clase del modelo que importa una clase de la vista",
         "**Dependencia en la direccion equivocada**: el modelo depende de la vista, de modo que no compila ni se "
         "prueba sin ella y cualquier cambio de interfaz lo afecta.",
         "Eliminar el import. Si el modelo necesita avisar, define su propia interfaz (`ObservadorInventario`, "
         "Ejercicio 4) y la vista o el controlador la implementan y se registran; la flecha vuelve a apuntar hacia "
         "el modelo."],
    ]
    tabla(doc, ["Fragmento", "Violacion", "Que componente hace el trabajo de otro / dependencia", "Correccion propuesta"],
          filas, [1.6, 3.4, 5.6, 5.9], tam=8.5)
    encabezado(doc, "Ejemplos de correccion", 2)
    parrafo(doc, "B. JOptionPane en el modelo:")
    codigo(doc, "// Antes (modelo)                                   // Despues (modelo)\n"
                "public void vender(String n, int c) {                public int vender(String n, int c)\n"
                "    if (c > p.getCantidad()) {                               throws StockInsuficienteException {\n"
                "        JOptionPane.showMessageDialog(null,              if (c > p.getCantidad())\n"
                "            \"Stock insuficiente\");                          throw new StockInsuficienteException(...);\n"
                "        return;                                          ...\n"
                "    } ...                                            }\n"
                "// Despues (controlador)\n"
                "try { vista.mostrarVentaOk(inventario.vender(n, c)); }\n"
                "catch (StockInsuficienteException e) { vista.mostrarStockInsuficiente(); }  // la vista usa JOptionPane")
    parrafo(doc, "D. Controlador que arma filas:")
    codigo(doc, "// Antes (controlador)\n"
                "for (Producto p : inventario.getProductos())\n"
                "    modeloTabla.addRow(new Object[]{p.getNombre(), p.getCantidad(), \"$\" + p.getPrecio()});\n"
                "// Despues (controlador)                        // Despues (vista)\n"
                "vista.mostrarListado(inventario.getProductos()); public void mostrarListado(List<Producto> ps) {\n"
                "                                                    modeloTabla.setRowCount(0);\n"
                "                                                    for (Producto p : ps) modeloTabla.addRow(...);\n"
                "                                                }")
    parrafo(doc, "E. Descuento por volumen en la vista:")
    codigo(doc, "// Antes (vista)\n"
                "double total = cantidad * precio;\n"
                "if (cantidad >= 100) total *= 0.90;          // regla de negocio en la interfaz\n"
                "lblTotal.setText(String.format(\"$%,.2f\", total));\n"
                "// Despues (modelo)                            // Despues (vista)\n"
                "public double calcularTotal() {                lblTotal.setText(String.format(\"$%,.2f\", total));\n"
                "    double t = cantidad * precio;\n"
                "    return cantidad >= 100 ? t * 0.90 : t;\n}")
    resultado(doc, "Tabla completa con las seis violaciones identificadas y su correccion.",
              "Seis violaciones identificadas indicando que componente hace el trabajo de cual (o que dependencia "
              "esta invertida), correccion propuesta y ejemplos de codigo antes/despues.")


# ============================================================================ Ejercicio 8

def ejercicio8(doc, resumen):
    encabezado(doc, "Ejercicio 8. Proyecto integrador: prestamos de biblioteca", 1)
    parrafo(doc, "Aplicacion de dominio pequeno y bien delimitado en `ejercicio8/`: una biblioteca registra libros y "
                 "socios, presta y recibe libros, cobra multas por retraso y simula el paso de los dias para "
                 "demostrarlo.")
    codigo(doc, "ejercicio8/src/\n"
                "  modelo/        Biblioteca (reglas), Libro, Socio, Prestamo, EstadoLibro, EstadoSocio,\n"
                "                 ObservadorBiblioteca, RepositorioBiblioteca, ReglaNegocioException y 5 subclases\n"
                "  vista/         VistaBiblioteca (interfaz), VistaConsola, VistaSwing, Formato\n"
                "  controlador/   ControladorBiblioteca (casos de uso comunes), ControladorConsola, ControladorSwing\n"
                "  persistencia/  RepositorioMemoria, RepositorioArchivo (libros.csv, socios.csv, prestamos.csv), Csv\n"
                "  soporte/       RelojAjustable (Clock que se puede adelantar para la demostracion)\n"
                "  Main.java      ensambla; java Main (consola + ventana) | --consola | --gui\n"
                "ejercicio8/test/ BibliotecaTest, RepositorioTest, ControladorBibliotecaTest, DosVistasTest")
    encabezado(doc, "Requisitos y como se cumplen", 2)
    ej8 = [r for r in resumen if r["ejercicio"] == 8]
    tabla(doc, ["Requisito", "Implementacion", "Verificacion"], [
        ["Modelo con al menos tres reglas no triviales cubiertas por pruebas",
         "R1 max. 3 prestamos activos por socio; R2 un libro prestado no se presta otra vez; R3 multa de $5 por dia "
         "de retraso (plazo 7 dias); R4 socio con multa no pide prestado y el pago debe ser > 0 y <= deuda.",
         "BibliotecaTest: 15 pruebas agrupadas por regla, con reloj fijo (incluye limites: dia 7 sin multa, dia 8 "
         "con $5)."],
        ["Dos vistas simultaneas sobre el mismo modelo con notificacion",
         "ObservadorBiblioteca en el modelo; consola y ventana Swing corren a la vez (`java Main`).",
         "DosVistasTest: un prestamo hecho en la consola aparece en la ventana y una devolucion hecha en la ventana "
         "aparece en la consola."],
        ["Persistencia intercambiable detras de una interfaz",
         "RepositorioBiblioteca (modelo) con RepositorioMemoria y RepositorioArchivo; se cambia una linea de Main.",
         "RepositorioTest (mismo escenario con ambas, y reapertura del archivo) + " + ej8[1]["obtenido"] + "."],
        ["Diagrama de clases por paquete y diagrama de secuencia",
         "`ejercicio8/diagramas/Ej8_Biblioteca.mdj`: diagrama de paquetes, 6 diagramas de clases (uno por paquete) "
         "y la secuencia \"Prestar un libro desde la ventana\".",
         "Validados (referencias e IDs) y abiertos por StarUML 6.3.4 en modo CLI (exporto 8/8 diagramas)."],
        ["Media cuartilla justificando donde va cada regla", "Seccion siguiente.", ""],
    ], [4, 7, 5.5], tam=8.5)

    texto = leer("evidencias/pruebas_ej8.txt")
    png = terminal_png(texto, EVID / "ej8_pruebas_verde.png", "java ... RunTests modelo persistencia controlador")
    imagen(doc, png, "Figura 11. Pruebas del proyecto integrador: 23 de 23 en verde.", Cm(16))

    encabezado(doc, "Evidencia de funcionamiento", 2)
    imagen(doc, RAIZ / "ejercicio8/evidencias/ej8_1_ventana_tras_prestamos_en_consola.png",
           "Figura 12. La consola presto L01 y L03 a Ana y L02 a Luis; la ventana lo refleja.")
    imagen(doc, RAIZ / "ejercicio8/evidencias/ej8_2_ventana_9_dias_despues_vencidos.png",
           "Figura 13. Tras avanzar 9 dias, el modelo marca los prestamos como vencidos y la vista los pinta de rojo.")
    imagen(doc, RAIZ / "ejercicio8/evidencias/ej8_3_ventana_multa_y_prestamo_rechazado.png",
           "Figura 14. Ana devuelve L01 dos dias tarde ($10 de multa) y la regla R4 rechaza su siguiente prestamo.")

    encabezado(doc, "Diagramas (StarUML)", 2)
    base = RAIZ / "ejercicio8/diagramas/vista_previa/Ej8_Biblioteca"
    imagen(doc, base / "Paquetes_-_Biblioteca.png", "Figura 15. Diagrama de paquetes: todas las dependencias "
                                                    "terminan en modelo.", Cm(11))
    imagen(doc, base / "Clases_-_paquete_modelo.png", "Figura 16. Clases del paquete modelo.")
    imagen(doc, base / "Clases_-_paquete_vista.png", "Figura 17. Clases del paquete vista.")
    imagen(doc, base / "Clases_-_paquete_controlador.png", "Figura 18. Clases del paquete controlador.")
    imagen(doc, base / "Clases_-_paquete_persistencia.png", "Figura 19. Clases del paquete persistencia (realizan "
                                                            "la interfaz del modelo).")
    imagen(doc, base / "Prestar_un_libro_desde_la_ventana.png",
           "Figura 20. Diagrama de secuencia de una operacion completa: prestar un libro desde la ventana, con la "
           "notificacion a ambas vistas.")

    encabezado(doc, "Justificacion: donde se coloco cada regla y por que", 2)
    parrafo(doc, "Las cuatro reglas (R1 a R4) estan en `modelo.Biblioteca`, porque son invariantes del negocio que "
                 "deben cumplirse sin importar desde donde llegue la operacion: la consola, la ventana o una prueba. "
                 "Cada operacion (`prestar`, `devolver`, `pagarMulta`) valida sus reglas dentro de un bloque "
                 "sincronizado y solo despues modifica el estado y avisa a los observadores, de modo que dos vistas "
                 "trabajando a la vez nunca ven un estado intermedio. El calculo de la multa es una funcion pura "
                 "(`calcularMulta(fechaLimite, fechaDevolucion)`) para probar sus limites sin montar escenarios, y la "
                 "fecha actual llega como `java.time.Clock` inyectado: asi el modelo es determinista en las pruebas y "
                 "la demostracion puede adelantar el tiempo con `RelojAjustable` sin que el modelo sepa que se trata "
                 "de una simulacion.")
    parrafo(doc, "Que un prestamo este vencido o que un socio pueda pedir libros tambien es conocimiento del negocio, "
                 "por eso lo calcula el modelo (`EstadoLibro`, `EstadoSocio`); la vista solo decide pintarlo de rojo "
                 "y como escribir fechas y montos (`vista.Formato`). Convertir el texto \"abc\" del monto o de los "
                 "dias a numero es trabajo del `ControladorBiblioteca`, compartido por ambas vistas: es la leccion "
                 "del Ejercicio 3, donde esa logica quedo duplicada. Por ultimo, el modelo define la interfaz "
                 "`RepositorioBiblioteca` en sus propios terminos y las implementaciones de memoria y archivo viven "
                 "en `persistencia`, de modo que ninguna regla depende de como se guardan los datos.")

    encabezado(doc, "Comparacion con MVP y MVVM", 2)
    tabla(doc, ["Aspecto", "MVC (este proyecto)", "MVP", "MVVM"], [
        ["Quien recibe la accion del usuario", "La vista la reenvia al controlador (listeners registrados por el "
         "controlador).", "La vista la reenvia al presentador.", "La vista ejecuta un comando del ViewModel por "
         "enlace (binding)."],
        ["Como se actualiza la vista", "Clasico: la vista observa al modelo. Aqui el controlador se suscribe y "
         "empuja datos a la vista.", "El presentador llama metodos de una interfaz de vista pasiva.",
         "La vista se enlaza a propiedades observables del ViewModel y se actualiza sola."],
        ["Conocimiento entre piezas", "Vista y controlador conocen al modelo; el modelo no conoce a nadie.",
         "La vista no conoce al modelo; el presentador conoce ambos (vista por interfaz).",
         "La vista no conoce al modelo; el ViewModel no conoce a la vista."],
        ["Pruebas", "Modelo con JUnit; controlador con una vista falsa.", "Presentador con una vista falsa "
         "(mock de la interfaz).", "ViewModel sin interfaz grafica (propiedades y comandos)."],
        ["Tecnologias tipicas", "Swing, Smalltalk, frameworks web (Spring MVC, Rails).",
         "Android clasico, WinForms, GWT.", "WPF/XAML, JavaFX con propiedades, Android Jetpack, Angular/Vue."],
    ], [3.2, 4.6, 4.2, 4.5], tam=8.5)
    parrafo(doc, "Observacion honesta sobre el diseno obtenido: como las vistas implementan `VistaBiblioteca` y el "
                 "controlador les empuja los datos ya calculados, la variante que se construyo esta muy cerca de "
                 "MVP con vista pasiva; lo que la mantiene en MVC es que el modelo publica sus cambios a observadores "
                 "y que ambas vistas se sincronizan por esos avisos y no porque un presentador las coordine. Pasar a "
                 "MVVM implicaria reemplazar los metodos `mostrar...` por propiedades observables (por ejemplo, "
                 "`ObservableList<EstadoLibro>` en JavaFX) enlazadas a la tabla.")
    resultado(doc, "Tres reglas no triviales con pruebas; dos vistas simultaneas con notificacion; persistencia "
                   "intercambiable; diagramas de clases por paquete y de secuencia; justificacion de media cuartilla.",
              f"{ej8[0]['obtenido']}; {ej8[1]['obtenido']}; 8 diagramas en StarUML; justificacion y comparacion "
              "MVC/MVP/MVVM.", all(r["ok"] for r in ej8))


# ============================================================================ documento

def portada(doc):
    for _ in range(3):
        doc.add_paragraph()
    parrafo(doc, "INSTITUTO TECNOLOGICO DE SONORA", negrita=True, tam=16, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    parrafo(doc, "Ingenieria en Software - Arquitectura de Software", tam=13, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    parrafo(doc, "Serie de ejercicios sobre el patron MVC", negrita=True, tam=22, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    parrafo(doc, "Respuestas, analisis y registro de resultados - Ejercicios 0 a 8 (Java)", tam=13,
            alinear=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(4):
        doc.add_paragraph()
    parrafo(doc, "Nombre del alumno: ______________________________________", tam=12, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    parrafo(doc, "Grupo: ____________", tam=12, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def introduccion(doc, resumen):
    encabezado(doc, "Como esta organizada la entrega", 1)
    parrafo(doc, "Todo el trabajo esta en un repositorio git con un commit y un tag por ejercicio, para que cada "
                 "entrega se pueda revisar y comparar con `git diff`. El PDF enumera los ejercicios 0 a 8 (nueve, "
                 "aunque el titulo dice \"los diez ejercicios\"); no existe un Ejercicio 9.")
    tabla(doc, ["Ruta", "Contenido"], [
        ["ejercicio0/SistemaInventario.java", "Codigo original, sin cambios (Ej. 0 y oraculo del Ej. 1)."],
        ["src/ (modelo, vista, controlador, persistencia, App)", "Inventario MVC; su estado en cada ejercicio esta en "
                                                                  "los tags ej1 ... ej6."],
        ["test/", "Pruebas JUnit 5 (modelo, controlador, persistencia), lanzador RunTests y generador de capturas."],
        ["ejercicio8/", "Proyecto integrador (codigo, pruebas, diagramas, evidencias)."],
        ["diagramas/*.mdj, ejercicio8/diagramas/*.mdj", "Diagramas en formato nativo de StarUML; imagenes/ (exportadas "
                                                        "por StarUML) y vista_previa/ (usadas en este documento)."],
        ["evidencias/", "Salidas reales de las verificaciones: pruebas, diffs, capturas, CSV."],
        ["scripts/", "verificar_todo.py (todas las verificaciones), verificar_ej1.py, verificar_ej6.py, "
                     "generar_mdj.py, exportar_staruml.py, render_mdj.py, generar_docx.py."],
        ["lib/", "JUnit 5.10.1 (Jupiter + Platform)."],
    ], [6, 10.5], tam=9)
    parrafo(doc, "Comandos (desde la carpeta del proyecto, en Windows):")
    codigo(doc, "javac -encoding UTF-8 -d out/app (todos los .java de src)\n"
                "java -cp out/app App                      # consola\n"
                "java -cp out/app App --dual               # consola + ventana sincronizadas\n"
                "python scripts/verificar_todo.py          # compila, prueba y verifica todos los ejercicios")
    encabezado(doc, "Resumen de verificacion", 2)
    tabla(doc, ["Ej.", "Resultado esperado", "Resultado obtenido", ""],
          [[str(r["ejercicio"]), r["esperado"], r["obtenido"], "OK" if r["ok"] else "FALLA"] for r in resumen],
          [1.3, 5.7, 8, 1.5], tam=8.5, colores=["E2EFDA" if r["ok"] else "F8CBAD" for r in resumen])
    tags = git("log", "--oneline", "--decorate", "--no-color")
    parrafo(doc, "Historial del repositorio:")
    codigo(doc, tags, tam=8)
    doc.add_page_break()


def anexos(doc):
    encabezado(doc, "Anexo A. Diagramas StarUML", 1)
    parrafo(doc, "Los archivos `.mdj` se generaron con `scripts/generar_mdj.py` a partir del codigo de cada tag (no "
                 "se dibujaron a mano), por lo que reflejan exactamente las clases, atributos, operaciones y "
                 "dependencias entregadas. Cada archivo se valido (identificadores unicos, todas las referencias "
                 "`$ref` resueltas, cada `_parent` coherente) y se abrio con StarUML 6.3.4 mediante su interfaz de "
                 "linea de comandos (`StarUML.exe image ...`, ventana oculta), que exporto todos los diagramas a "
                 "`imagenes/`. Esas exportaciones llevan la marca de agua \"UNREGISTERED\" de la copia sin licencia; "
                 "por eso las figuras de este documento son vistas previas dibujadas a partir del mismo `.mdj` "
                 "(`scripts/render_mdj.py`).")
    codigo(doc, leer("evidencias/diagramas_validacion.txt"), tam=7.5)
    tabla(doc, ["Archivo", "Diagramas"], [
        ["diagramas/Ej1_Paquetes_MVC.mdj", "Paquetes - Ejercicio 1; Clases - Ejercicio 1"],
        ["diagramas/Ej3_Dos_Vistas.mdj", "Paquetes - Ejercicio 3; Clases - Ejercicio 3"],
        ["diagramas/Ej6_Persistencia_DAO.mdj", "Paquetes - Ejercicio 6; Clases - Ejercicio 6"],
        ["ejercicio8/diagramas/Ej8_Biblioteca.mdj", "Paquetes; Clases por paquete (modelo, vista, controlador, "
                                                    "persistencia, soporte, raiz); Secuencia: prestar un libro"],
    ], [7, 9.5], tam=9)
    encabezado(doc, "Anexo B. Publicacion en GitHub (pendiente)", 1)
    parrafo(doc, "El Ejercicio 1 pide subir el repositorio y agregar al profesor como colaborador. El repositorio local "
                 "esta listo; falta crear el repositorio remoto con la cuenta del alumno y ejecutar:")
    codigo(doc, "git remote add origin https://github.com/<usuario>/<repositorio>.git\n"
                "git push -u origin main --tags\n"
                "# GitHub > Settings > Collaborators > Add people > copequintanapotro")


def main():
    resumen = json.loads(leer("evidencias/resumen_verificacion.json"))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.orientation = WD_ORIENT.PORTRAIT
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(2.5))
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for nivel, tam, color in ((1, 16, "2B579A"), (2, 13, "2B579A")):
        st = doc.styles[f"Heading {nivel}"]
        st.font.size = Pt(tam)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.name = "Calibri"
    pie = sec.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = pie.add_run("Serie de ejercicios MVC - Arquitectura de Software - pagina ")
    r.font.size = Pt(8)
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            pie.add_run()._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = texto
            pie.add_run()._r.append(it)

    portada(doc)
    introduccion(doc, resumen)
    for funcion in (ejercicio0, ejercicio1, ejercicio2, ejercicio3, ejercicio4, ejercicio5, ejercicio6, ejercicio7,
                    ejercicio8):
        if funcion in (ejercicio0, ejercicio5, ejercicio7):
            funcion(doc)
        else:
            funcion(doc, resumen)
        doc.add_page_break()
    anexos(doc)
    doc.save(SALIDA)
    print("Documento:", SALIDA)


if __name__ == "__main__":
    main()
